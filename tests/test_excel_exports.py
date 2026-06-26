"""Regression tests for Excel export structural integrity and shared styling.

These exports previously had no direct test coverage. They are exercised with
empty input (which every service tolerates, producing an empty-but-valid
workbook) so a structural break surfaces here, and the shared style palette is
asserted against webapp.app_services.excel_styles.
"""

import importlib

import pytest

pytest.importorskip("openpyxl")
from openpyxl import load_workbook  # noqa: E402

from webapp.app_services import excel_styles  # noqa: E402

EXPORT_SERVICES = [
    "aging_dashboard",
    "bank_recon_status",
    "budget_actual",
    "depreciation_calc",
    "fuel_tax_credits",
    "payg_instalment",
    "payg_reconciliation",
    "payroll_tax",
    "prepayment_tracker",
    "stp_tracker",
]


@pytest.mark.parametrize("service_name", EXPORT_SERVICES)
def test_export_to_excel_produces_valid_workbook(service_name):
    module = importlib.import_module(f"webapp.app_services.{service_name}_service")
    workbook = load_workbook(module.export_to_excel({}))
    assert workbook.sheetnames  # at least one sheet, file opens cleanly


def test_sanitize_workbook_neutralizes_formula_cells():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=1+2"
    ws["A2"] = "+cmd"
    ws["A3"] = "-2+3"
    ws["A4"] = "@SUM(1,2)"
    ws["A5"] = "Normal Name"
    ws["A6"] = 1234

    excel_styles.sanitize_workbook(wb)

    for ref in ("A1", "A2", "A3", "A4"):
        assert ws[ref].data_type == "s", f"{ref} should be text, not a formula"
        assert ws[ref].value.startswith("'")
    assert ws["A5"].value == "Normal Name"  # untouched
    assert ws["A6"].value == 1234  # numeric untouched


def test_export_neutralizes_formula_injection_end_to_end():
    """A malicious Xero contact/account name must not become a live formula."""
    from webapp.app_services import bank_recon_status_service as svc

    data = {
        "data": {
            "accounts": [
                {
                    "name": '=HYPERLINK("http://attacker.example","x")',
                    "code": "+cmd|'/C calc'!A0",
                    "status": "ok",
                    "unreconciled_items": [
                        {
                            "date": "2025-01-01",
                            "type": "SPEND",
                            "contact": "@SUM(1,2)",
                            "reference": "-2+2",
                            "amount": 100,
                        }
                    ],
                }
            ]
        }
    }
    wb = load_workbook(svc.export_to_excel(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                # No cell may be a live formula, and any injected string must
                # have been neutralized with a leading quote.
                assert cell.data_type != "f", f"formula leaked at {cell.coordinate}"
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                    raise AssertionError(
                        f"unsanitized formula char at {cell.coordinate}: {cell.value!r}"
                    )


def test_style_factories_use_brand_palette():
    assert str(excel_styles.header_fill().fgColor.rgb).endswith(excel_styles.HEADER_COLOR)
    assert str(excel_styles.ok_fill().fgColor.rgb).endswith(excel_styles.OK_COLOR)
    assert str(excel_styles.warning_fill().fgColor.rgb).endswith(
        excel_styles.WARNING_COLOR
    )
    assert str(excel_styles.error_fill().fgColor.rgb).endswith(excel_styles.ERROR_COLOR)

    font = excel_styles.header_font()
    assert font.bold is True
    assert str(font.color.rgb).endswith(excel_styles.HEADER_FONT_COLOR)
